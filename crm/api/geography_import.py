import re
import unicodedata
from collections import defaultdict

import frappe
from frappe import _


HEADER_MAP = {
	"matinhtp": "province_code",
	"matinh": "province_code",
	"matinhthanhpho": "province_code",
	"tentinhtp": "province_name",
	"tentinh": "province_name",
	"tentinhthanhpho": "province_name",
	"maxaphuong": "ward_code",
	"maxa": "ward_code",
	"maphuong": "ward_code",
	"maxaphuongthitran": "ward_code",
	"tenxaphuong": "ward_name",
	"tenxa": "ward_name",
	"tenphuong": "ward_name",
	"tenxaphuongthitran": "ward_name",
	"matruong": "school_code",
	"tentruong": "school_name",
	"diachi": "address",
	"khuvuc": "region_code",
}

REQUIRED_COLUMNS = {
	"province_code",
	"province_name",
	"ward_code",
	"ward_name",
	"school_code",
	"school_name",
	"address",
	"region_code",
}


@frappe.whitelist()
def import_geography_high_schools(file_url: str):
	frappe.only_for("System Manager", True)

	if not file_url:
		frappe.throw(_("File URL is required"))

	rows = _read_excel_rows(file_url)
	if not rows:
		frappe.throw(_("The uploaded file is empty"))

	headers = _get_headers(rows[0])
	missing = sorted(REQUIRED_COLUMNS - set(headers.values()))
	if missing:
		frappe.throw(_("Missing required columns: {0}").format(", ".join(missing)))

	summary = {
		"processed": 0,
		"created": defaultdict(int),
		"updated": defaultdict(int),
		"skipped": 0,
		"errors": [],
	}

	for row_index, row in enumerate(rows[1:], start=2):
		data = _row_to_dict(headers, row)
		if not any(data.values()):
			continue

		savepoint = f"crm_geography_import_{row_index}"
		frappe.db.savepoint(savepoint)
		try:
			_import_row(data)
			summary["processed"] += 1
			for key, value in frappe.flags.crm_geography_import_counts["created"].items():
				summary["created"][key] += value
			for key, value in frappe.flags.crm_geography_import_counts["updated"].items():
				summary["updated"][key] += value
		except Exception as exc:
			frappe.db.rollback(save_point=savepoint)
			summary["skipped"] += 1
			summary["errors"].append({"row": row_index, "message": str(exc)})

	frappe.db.commit()
	summary["created"] = dict(summary["created"])
	summary["updated"] = dict(summary["updated"])
	return summary


def _read_excel_rows(file_url):
	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		frappe.throw(_("Uploaded file not found"))

	file_doc = frappe.get_doc("File", file_name)
	content = file_doc.get_content()

	try:
		from openpyxl import load_workbook
	except ImportError:
		frappe.throw(_("openpyxl is required to import Excel files"))

	from io import BytesIO

	workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
	sheet = workbook.active
	rows = []
	for row in sheet.iter_rows():
		rows.append([_cell_value(cell) for cell in row])
	return rows


def _cell_value(cell):
	value = cell.value
	if value is None:
		return ""

	if isinstance(value, float) and value.is_integer():
		value = int(value)

	if isinstance(value, int):
		number_format = (cell.number_format or "").strip()
		if re.fullmatch(r"0+", number_format):
			return str(value).zfill(len(number_format))

	return str(value).strip()


def _get_headers(row):
	headers = {}
	for index, value in enumerate(row):
		key = HEADER_MAP.get(_normalize_header(value))
		if key:
			headers[index] = key
	return headers


def _normalize_header(value):
	value = unicodedata.normalize("NFD", str(value or ""))
	value = "".join(char for char in value if unicodedata.category(char) != "Mn")
	return re.sub(r"[^a-zA-Z0-9]", "", value).lower()


def _row_to_dict(headers, row):
	data = {}
	for index, key in headers.items():
		data[key] = row[index].strip() if index < len(row) else ""
	return data


def _import_row(data):
	frappe.flags.crm_geography_import_counts = {
		"created": defaultdict(int),
		"updated": defaultdict(int),
	}

	_required(data, "province_code", "Mã Tỉnh/TP")
	_required(data, "province_name", "Tên Tỉnh/TP")
	_required(data, "ward_code", "Mã Xã/Phường")
	_required(data, "ward_name", "Tên Xã/Phường")
	_required(data, "school_code", "Mã Trường")
	_required(data, "school_name", "Tên Trường")
	_required(data, "region_code", "Khu Vực")

	region = _upsert_doc(
		"CRM Region",
		{"region_code": data["region_code"]},
		{
			"region_name": data["region_code"],
			"region_code": data["region_code"],
		},
		fallback_filters={"region_name": data["region_code"]},
	)
	province = _upsert_doc(
		"CRM Province",
		{"province_code": data["province_code"]},
		{
			"province_name": data["province_name"],
			"province_code": data["province_code"],
			"city_type": "Province",
			"region": region.name,
		},
		fallback_filters={"province_name": data["province_name"]},
	)
	ward = _upsert_doc(
		"CRM Ward",
		{"ward_code": data["ward_code"], "province": province.name},
		{
			"ward_name": data["ward_name"],
			"ward_code": data["ward_code"],
			"province": province.name,
			"province_name": data["province_name"],
			"ward_type": "Ward",
		},
		fallback_filters={"ward_name": data["ward_name"], "province": province.name},
	)
	_upsert_doc(
		"CRM High School",
		{"school_code": data["school_code"], "province_code": data["province_code"]},
		{
			"school_name": data["school_name"],
			"school_code": data["school_code"],
			"province_code": data["province_code"],
			"province_name": data["province_name"],
			"ward_code": data["ward_code"],
			"ward_name": data["ward_name"],
			"address": data.get("address"),
		},
		fallback_filters={
			"school_name": data["school_name"],
			"ward_code": data["ward_code"],
			"province_code": data["province_code"],
		},
	)


def _required(data, fieldname, label):
	if not data.get(fieldname):
		frappe.throw(_("{0} is required").format(label))


def _upsert_doc(doctype, filters, values, fallback_filters=None):
	name = frappe.db.get_value(doctype, filters, "name")
	if not name and fallback_filters:
		name = frappe.db.get_value(doctype, fallback_filters, "name")

	if name:
		doc = frappe.get_doc(doctype, name)
		for key, new_value in values.items():
			if new_value is not None and doc.get(key) != new_value:
				doc.set(key, new_value)
		doc.save(ignore_permissions=True)
		frappe.flags.crm_geography_import_counts["updated"][doctype] += 1
		return doc

	doc = frappe.get_doc({"doctype": doctype, **values})
	doc.insert(ignore_permissions=True)
	frappe.flags.crm_geography_import_counts["created"][doctype] += 1
	return doc
