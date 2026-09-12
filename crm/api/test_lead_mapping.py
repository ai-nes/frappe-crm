import json
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase

from crm.api import lead_mapping
from crm.api.lead_mapping import (
	LeadMappingError,
	_normalize_public_lead_payload,
	_parse_csv_rows,
	_parse_import_file,
	_parse_public_payload,
	_resolve_assignment,
	create_public_lead,
	get_public_high_schools,
	get_public_leads,
	get_public_majors,
	get_public_provinces,
	get_public_wards,
	split_multi_value,
)


def _quick_import_row(**overrides):
	row = {
		"student_name": "Nguyễn Văn An",
		"phone": "0900000000",
		"province": "Ho Chi Minh City",
		"high_school": "High School 1",
		"source": "Promoter",
		"assigned_to": None,
	}
	row.update(overrides)
	return row


class TestLeadMappingContract(TestCase):
	def test_inspect_csv_preserves_arbitrary_headers_duplicate_labels_and_samples(self):
		upload = SimpleNamespace(
			filename="arbitrary.csv",
			read=lambda: b"\nName,Phone,Phone,Unknown\nAn,0900000000,0900000001,extra\n",
		)
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
		):
			result = lead_mapping.inspect_lead_import()

		self.assertEqual(
			result["headers"],
			[
				{"sourceIndex": 0, "label": "Name", "inferredField": None, "enabled": False},
				{"sourceIndex": 1, "label": "Phone", "inferredField": "phone", "enabled": True},
				{"sourceIndex": 2, "label": "Phone", "inferredField": "phone", "enabled": True},
				{"sourceIndex": 3, "label": "Unknown", "inferredField": None, "enabled": False},
			],
		)
		self.assertEqual(
			result["sampleRows"], [{"row": 3, "values": ["An", "0900000000", "0900000001", "extra"]}]
		)
		self.assertEqual(
			result["requiredFields"], ["student_name", "phone", "province", "high_school", "source"]
		)

	def test_inspect_xlsx_returns_json_safe_samples_and_physical_rows(self):
		from openpyxl import Workbook

		workbook = Workbook()
		worksheet = workbook.active
		worksheet.append([])
		worksheet.append(["Họ và tên", "Ngày", "Điểm"])
		worksheet.append(["An", date(2026, 9, 9), 9.5])
		worksheet.append(["B", None, 10])
		content = BytesIO()
		workbook.save(content)
		upload = SimpleNamespace(filename="arbitrary.xlsx", read=lambda: content.getvalue())

		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
		):
			result = lead_mapping.inspect_lead_import()

		self.assertEqual(result["sampleRows"][0], {"row": 3, "values": ["An", "2026-09-09", "9.5"]})
		self.assertEqual(result["sampleRows"][1], {"row": 4, "values": ["B", None, "10"]})
		self.assertLessEqual(len(result["sampleRows"]), lead_mapping.MAX_IMPORT_SAMPLE_ROWS)

	def test_inspect_rejects_without_lead_create_permission_before_reading_file(self):
		upload = SimpleNamespace(filename="arbitrary.csv", read=lambda: b"Name\nAn\n")
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=False),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
			patch.object(lead_mapping, "_uploaded_import_file") as uploaded_file,
			self.assertRaises(LeadMappingError) as context,
		):
			lead_mapping.inspect_lead_import()

		self.assertEqual(context.exception.code, "FORBIDDEN")
		uploaded_file.assert_not_called()

	def test_mapped_preview_requires_unique_required_targets_and_rejects_server_fields(self):
		upload = SimpleNamespace(filename="arbitrary.csv", read=lambda: b"Name,Other\nAn,value\n")
		base_form = {
			"column_mapping": json.dumps([{"sourceIndex": 0, "targetField": "student_name", "enabled": True}])
		}
		for mapping, expected_code in (
			(base_form["column_mapping"], "MISSING_REQUIRED_MAPPING"),
			(
				json.dumps(
					[
						{"sourceIndex": 0, "targetField": "student_name", "enabled": True},
						{"sourceIndex": 1, "targetField": "student_name", "enabled": True},
					]
				),
				"DUPLICATE_TARGET_FIELD",
			),
			(
				json.dumps([{"sourceIndex": 0, "targetField": "campaign", "enabled": True}]),
				"SERVER_MANAGED_FIELD",
			),
		):
			with self.subTest(expected_code=expected_code):
				with (
					patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
					patch.object(lead_mapping.frappe, "has_permission", return_value=True),
					patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
					patch.object(
						lead_mapping.frappe,
						"form_dict",
						{**base_form, "column_mapping": mapping},
						create=True,
					),
				):
					with self.assertRaises(LeadMappingError) as context:
						lead_mapping.preview_lead_import.__wrapped__()

				self.assertEqual(context.exception.code, expected_code)

	def test_mapped_preview_uses_source_indexes_and_ignores_disabled_columns(self):
		upload = SimpleNamespace(
			filename="arbitrary.csv",
			read=lambda: b"Name,Phone,Province,School,Source,Ignore\n"
			b"An,0900000000,HCM,School 1,Promoter,not imported\n",
		)
		mapping = [
			{"sourceIndex": 0, "targetField": "student_name", "enabled": True},
			{"sourceIndex": 1, "targetField": "phone", "enabled": True},
			{"sourceIndex": 2, "targetField": "province", "enabled": True},
			{"sourceIndex": 3, "targetField": "high_school", "enabled": True},
			{"sourceIndex": 4, "targetField": "source", "enabled": True},
			{"sourceIndex": 5, "targetField": "email", "enabled": False},
		]
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
			patch.object(
				lead_mapping.frappe,
				"form_dict",
				{"column_mapping": json.dumps(mapping)},
				create=True,
			),
			patch.object(
				lead_mapping,
				"_normalize_lead_payload",
				return_value=(
					{"student_name": "An", "phone": "0900000000", "province": "PROVINCE-1"},
					[],
					[],
					None,
				),
			) as normalize_payload,
		):
			result = lead_mapping.preview_lead_import.__wrapped__()

		self.assertEqual(result["rows"][0]["row"], 2)
		self.assertEqual(
			result["mappedFields"], ["student_name", "phone", "province", "high_school", "source"]
		)
		self.assertEqual(result["ignoredColumns"], [{"sourceIndex": 5, "label": "Ignore"}])
		normalize_payload.assert_called_once_with(
			{
				"student_name": "An",
				"phone": "0900000000",
				"province": "HCM",
				"high_school": "School 1",
				"source": "Promoter",
			},
			allow_unassigned=True,
			require_import_fields=True,
			campaign_name=None,
		)

	def test_quick_multipart_commit_reparses_original_file_and_mapping(self):
		upload = SimpleNamespace(
			filename="arbitrary.csv",
			read=lambda: b"Name,Phone,Province,School,Source\nAn,0900000000,HCM,School 1,Promoter\n",
		)
		mapping = [
			{"sourceIndex": 0, "targetField": "student_name", "enabled": True},
			{"sourceIndex": 1, "targetField": "phone", "enabled": True},
			{"sourceIndex": 2, "targetField": "province", "enabled": True},
			{"sourceIndex": 3, "targetField": "high_school", "enabled": True},
			{"sourceIndex": 4, "targetField": "source", "enabled": True},
		]
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
			patch.object(
				lead_mapping.frappe, "form_dict", {"column_mapping": json.dumps(mapping)}, create=True
			),
			patch.object(lead_mapping, "_resolve_quick_import_campaign", return_value="Campaign 1"),
			patch.object(lead_mapping, "_parse_import_rows") as legacy_parser,
			patch.object(lead_mapping, "_create_lead", return_value={"name": "LEAD-1"}) as create_lead,
		):
			result = lead_mapping.import_leads.__wrapped__(
				import_mode="quick_create", campaign_code="CAM-2026-00001"
			)

		self.assertEqual(result["created"], 1)
		self.assertEqual(result["filename"], "arbitrary.csv")
		legacy_parser.assert_not_called()
		create_lead.assert_called_once_with(
			{
				"student_name": "An",
				"phone": "0900000000",
				"province": "HCM",
				"high_school": "School 1",
				"source": "Promoter",
			},
			allow_unassigned=True,
			require_import_fields=True,
			campaign_name="Campaign 1",
		)

	def test_parse_import_file_accepts_template_xlsx_headers(self):
		from openpyxl import Workbook

		workbook = Workbook()
		worksheet = workbook.active
		worksheet.append(["Họ và tên", "Di động", "Tỉnh/Thành Phố", "Trường THPT", "Nguồn", "Ngành quan tâm"])
		worksheet.append(
			[
				"Nguyễn Văn An",
				"0900000000",
				"Hồ Chí Minh",
				"THPT Nguyễn Huệ",
				"Promoter",
				"Kỹ thuật phần mềm",
			]
		)
		content = BytesIO()
		workbook.save(content)

		rows = _parse_import_file(content.getvalue(), "lead-template.xlsx")

		self.assertEqual(rows[0]["student_name"], "Nguyễn Văn An")
		self.assertEqual(rows[0]["major"], "Kỹ thuật phần mềm")

	def test_parse_import_file_accepts_formatted_template_rows_and_ignores_status(self):
		from openpyxl import Workbook

		workbook = Workbook()
		worksheet = workbook.active
		worksheet.append([])
		worksheet.append(["Mẫu import Lead"])
		worksheet.append(["Xóa dòng ví dụ trước khi import"])
		worksheet.append(
			[
				"Họ và tên",
				"Di động",
				"Tỉnh/Thành Phố",
				"Trường THPT",
				"Tình trạng Lead",
				"Nguồn",
				"Giao cho",
				"Nguyện vọng vào FPT",
				"Mô tả",
				"Năm tuyển sinh",
				"Ngành quan tâm",
			]
		)
		worksheet.append(
			[
				"Nguyễn Văn An",
				"0900000000",
				"Hồ Chí Minh",
				"THPT Nguyễn Huệ",
				"Mới",
				"Promoter",
				None,
				"Kỹ thuật phần mềm",
				"Quan tâm học bổng",
				2026,
				"Kỹ thuật phần mềm",
			]
		)
		content = BytesIO()
		workbook.save(content)

		rows = _parse_import_file(content.getvalue(), "lead-template.xlsx")

		self.assertEqual(rows[0]["student_name"], "Nguyễn Văn An")
		self.assertEqual(rows[0]["description"], "Quan tâm học bổng")
		self.assertNotIn("enrollment_status", rows[0])

	def test_parse_import_file_allows_optional_major_for_quick_import_file(self):
		rows = _parse_import_file(
			"Họ và tên,Di động,Tỉnh/Thành Phố,Trường THPT,Nguồn\n"
			"An,0900000000,Hà Nội,THPT Chu Văn An,Promoter",
			"leads.csv",
		)

		self.assertEqual(rows[0]["student_name"], "An")
		self.assertNotIn("major", rows[0])

	def test_parse_import_file_rejects_server_managed_campaign_column(self):
		with self.assertRaises(LeadMappingError) as context:
			_parse_import_file(
				"Họ và tên,Di động,Tỉnh/Thành Phố,Trường THPT,Nguồn,Campaign\n"
				"An,0900000000,Hà Nội,THPT Chu Văn An,Promoter,CAM-2026-00001",
				"leads.csv",
			)

		self.assertEqual(context.exception.code, "UNKNOWN_FIELD")

	def test_parse_payload_accepts_religion(self):
		payload = lead_mapping._parse_payload({"student_name": "An", "religion": "Phật giáo"})

		self.assertEqual(payload["religion"], "Phật giáo")

	def test_uploaded_import_file_reads_at_most_configured_limit(self):
		read_sizes = []
		stream = SimpleNamespace(
			read=lambda size: (
				read_sizes.append(size),
				b"x" * (lead_mapping.MAX_IMPORT_FILE_BYTES + 1),
			)[1]
		)
		upload = SimpleNamespace(filename="leads.csv", stream=stream)
		with (
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
			self.assertRaises(LeadMappingError) as context,
		):
			lead_mapping._uploaded_import_file()

		self.assertEqual(context.exception.code, "FILE_TOO_LARGE")
		self.assertEqual(read_sizes, [lead_mapping.MAX_IMPORT_FILE_BYTES + 1])

	def test_normalize_quick_import_allows_missing_optional_fields(self):
		fake_frappe = SimpleNamespace(
			db=SimpleNamespace(exists=lambda *_args, **_kwargs: True),
			as_json=lambda value: value,
		)
		with (
			patch.object(lead_mapping, "frappe", fake_frappe),
			patch.object(lead_mapping, "_resolve_province", return_value="PROVINCE-1"),
			patch.object(lead_mapping, "_resolve_source", return_value="Promoter"),
			patch.object(lead_mapping, "resolve_high_school_strict", return_value="SCHOOL-1"),
			patch.object(lead_mapping, "_resolve_assignment", return_value=(None, None, None)),
			patch.object(lead_mapping, "_current_admission_year", return_value=None),
		):
			values, _tags, _events, _assigned_user = lead_mapping._normalize_lead_payload(
				{
					"student_name": "An",
					"phone": "0900000000",
					"religion": "Phật giáo",
					"province": "Hà Nội",
					"high_school": "THPT Chu Văn An",
					"source": "Promoter",
				},
				allow_unassigned=True,
				require_import_fields=True,
			)

		self.assertIsNone(values["major"])
		self.assertIsNone(values["aspiration"])
		self.assertIsNone(values["admission_year"])
		self.assertEqual(values["religion"], "Phật giáo")

	def test_resolve_assignment_allows_blank_owner_only_for_quick_import(self):
		self.assertEqual(_resolve_assignment(None, allow_unassigned=True), (None, None, None))

		with patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")):
			with self.assertRaises(LeadMappingError) as context:
				_resolve_assignment(None)

		self.assertEqual(context.exception.code, "ASSIGNED_TO_REQUIRED")

	def test_quick_import_owner_must_be_a_staff_name(self):
		fake_frappe = SimpleNamespace(
			session=SimpleNamespace(user="Administrator"),
			db=SimpleNamespace(
				get_value=lambda _doctype, _filters, *_args, **_kwargs: {
					"name": "CTV Sale",
					"user": "ctv@example.com",
					"is_active": 1,
					"campus": "HCM",
				}
				if isinstance(_filters, dict)
				else None
			),
		)
		with patch.object(lead_mapping, "frappe", fake_frappe):
			with self.assertRaises(LeadMappingError) as context:
				_resolve_assignment("ctv@example.com", require_staff_name=True)

		self.assertEqual(context.exception.code, "ASSIGNED_TO_REQUIRED")

	def test_preview_import_does_not_create_leads(self):
		upload = SimpleNamespace(
			filename="leads.csv",
			read=lambda: (
				"Họ và tên,Di động,Tỉnh/Thành Phố,Trường THPT,Nguồn,Ngành quan tâm\n"
				"Nguyễn Văn An,0900000000,Hồ Chí Minh,THPT Nguyễn Huệ,Promoter,Kỹ thuật phần mềm"
			).encode(),
		)
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
			patch.object(lead_mapping.frappe, "form_dict", {}, create=True),
			patch.object(
				lead_mapping,
				"_normalize_lead_payload",
				return_value=(
					{
						"student_name": "Nguyễn Văn An",
						"phone": "0900000000",
						"province": "Ho Chi Minh City",
						"high_school": "High School 1",
						"source": "Promoter",
						"major": "Software Engineering",
						"assigned_to": None,
					},
					[],
					[],
					None,
				),
			),
			patch.object(lead_mapping, "_create_lead") as create_lead,
		):
			result = lead_mapping.preview_lead_import.__wrapped__()

		self.assertEqual(result["valid"], 1)
		self.assertEqual(result["failed"], 0)
		self.assertIsNone(result["rows"][0]["fields"]["assigned_to"])
		self.assertNotIn("campaign", result["rows"][0]["fields"])
		create_lead.assert_not_called()

	def test_quick_import_creates_direct_leads_without_routing_blank_owner(self):
		row = {
			"student_name": "Nguyễn Văn An",
			"phone": "0900000000",
			"province": "Ho Chi Minh City",
			"high_school": "High School 1",
			"source": "Promoter",
			"major": "Software Engineering",
			"assigned_to": None,
		}
		fake_frappe = SimpleNamespace(
			session=SimpleNamespace(user="Administrator"),
			has_permission=lambda *_args, **_kwargs: True,
			db=SimpleNamespace(
				get_value=lambda *_args, **_kwargs: {"name": "Campaign 1", "status": "ACTIVE"},
				savepoint=lambda _name: None,
			),
		)
		with (
			patch.object(lead_mapping, "frappe", fake_frappe),
			patch.object(lead_mapping, "_parse_import_rows", return_value=[row]),
			patch.object(lead_mapping, "_create_lead", return_value={"name": "LEAD-1"}) as create_lead,
			patch.object(lead_mapping, "route_new_lead") as route,
		):
			result = lead_mapping.import_leads.__wrapped__(
				rows=[row], import_mode="quick_create", campaign_code="CAM-2026-00001"
			)

		self.assertEqual(result["created"], 1)
		create_lead.assert_called_once_with(
			row,
			allow_unassigned=True,
			require_import_fields=True,
			campaign_name="Campaign 1",
		)
		route.assert_not_called()

	def test_split_multi_value_deduplicates_csv_values(self):
		self.assertEqual(
			split_multi_value("Segment A; Segment B\nSegment A"),
			["Segment A", "Segment B"],
		)

	def test_parse_csv_rows_accepts_mapping_headers(self):
		rows = _parse_csv_rows(
			"\ufeffHọ và Tên,Di động,Tỉnh/Thành Phố,Nguồn,Giao cho\n"
			"Nguyễn Văn An,0900000000,Hồ Chí Minh,Promoter,sales@example.com"
		)

		self.assertEqual(
			rows,
			[
				{
					"student_name": "Nguyễn Văn An",
					"phone": "0900000000",
					"province": "Hồ Chí Minh",
					"source": "Promoter",
					"assigned_to": "sales@example.com",
				}
			],
		)

	def test_parse_csv_rows_requires_all_non_nullable_headers(self):
		with self.assertRaises(LeadMappingError) as context:
			_parse_csv_rows("Họ và Tên,Di động,Tỉnh/Thành Phố,Nguồn\nAn,0900000000,Hà Nội,Promoter")

		self.assertEqual(context.exception.code, "CSV_MISSING_HEADERS")
		self.assertIn("assigned_to", str(context.exception))

	def test_parse_csv_rows_accepts_campaign_header_without_source(self):
		rows = _parse_csv_rows(
			"Họ và Tên,Di động,Tỉnh/Thành Phố,Campaign,Giao cho\n"
			"Nguyễn Văn An,0900000000,Hồ Chí Minh,Campaign 1,sales@example.com"
		)

		self.assertEqual(rows[0]["campaign"], "Campaign 1")
		self.assertNotIn("source", rows[0])

	def test_public_payload_rejects_server_managed_fields(self):
		with self.assertRaises(LeadMappingError) as context:
			_parse_public_payload({"student_name": "An", "processing_status": "PROCESSED"})

		self.assertEqual(context.exception.code, "SERVER_MANAGED_FIELD")
		self.assertIn("processing_status", str(context.exception))

	def test_public_payload_requires_campaign_code_during_normalization(self):
		from crm.api.lead_mapping import _normalize_public_lead_payload

		with self.assertRaises(LeadMappingError) as context:
			_normalize_public_lead_payload({"student_name": "An"})

		self.assertEqual(context.exception.code, "REQUIRED_FIELD")
		self.assertIn("campaign_code", str(context.exception))

	def test_public_payload_normalizes_optional_intake_fields(self):
		with (
			patch("crm.api.lead_mapping._check_unique") as check_unique,
			patch("crm.api.lead_mapping._resolve_campaign_code", return_value="_Test Campaign"),
			patch("crm.api.lead_mapping.frappe.db.get_value", return_value=None),
		):
			values = _normalize_public_lead_payload(
				{
					"student_name": "An",
					"phone": "+84981000099",
					"id_number": "012-345 678 901",
					"campaign_code": "CAM-2026-00001",
					"segments": ["Scholarship", "Scholarship"],
					"assignment_priority": "HIGH",
				}
			)

		self.assertEqual(values["phone"], "0981000099")
		self.assertEqual(values["id_number"], "012345678901")
		self.assertEqual(frappe.parse_json(values["segments"]), ["Scholarship"])
		self.assertEqual(values["assignment_priority"], "high")
		self.assertEqual(values["campaign"], "_Test Campaign")
		check_unique.assert_not_called()

	def test_public_payload_keeps_repeat_submissions_for_one_student(self):
		with (
			patch("crm.api.lead_mapping._resolve_campaign_code", return_value="_Test Campaign"),
			patch("crm.api.lead_mapping.frappe.db.get_value", return_value=None),
		):
			first = _normalize_public_lead_payload(
				{
					"student_name": "An",
					"phone": "0981000099",
					"email": "an@example.com",
					"id_number": "012345678901",
					"campaign_code": "CAM-2026-00001",
				}
			)
			second = _normalize_public_lead_payload(
				{
					"student_name": "An",
					"phone": "0981000099",
					"email": "an@example.com",
					"id_number": "012345678901",
					"campaign_code": "CAM-2026-00001",
				}
			)

		self.assertEqual(first["id_number"], second["id_number"])
		self.assertEqual(first["phone"], second["phone"])
		self.assertEqual(first["email"], second["email"])

	def test_public_payload_rejects_removed_cccd_field(self):
		with self.assertRaises(LeadMappingError) as context:
			_parse_public_payload(
				{
					"student_name": "An",
					"campaign_code": "CAM-2026-00001",
					"cccd": "012345678901",
				}
			)

		self.assertEqual(context.exception.code, "UNKNOWN_FIELD")
		self.assertIn("cccd", str(context.exception))

	@patch("crm.api.lead_mapping._create_public_lead", return_value={"ok": True})
	def test_public_endpoint_ignores_frappe_cmd_metadata(self, create_lead):
		result = create_public_lead(student_name="An", cmd="crm.api.lead_mapping.create_public_lead")

		self.assertEqual(result, {"ok": True})
		create_lead.assert_called_once_with({"student_name": "An"})

	@patch(
		"crm.api.lead_mapping.frappe.get_all",
		return_value=[{"name": "Ho Chi Minh City", "province_name": "Hồ Chí Minh", "province_code": "HCM"}],
	)
	def test_public_provinces_return_dropdown_items(self, get_all):
		result = get_public_provinces()

		self.assertEqual(
			result,
			{
				"items": [{"value": "Ho Chi Minh City", "label": "Hồ Chí Minh", "code": "HCM"}],
				"total": 1,
			},
		)
		get_all.assert_called_once()

	@patch("crm.api.lead_mapping.resolve_province", return_value="Ho Chi Minh City")
	@patch("crm.api.lead_mapping.frappe.db.exists", return_value=True)
	@patch(
		"crm.api.lead_mapping.frappe.get_all",
		return_value=[
			{
				"name": "Ward 1 - Ho Chi Minh City",
				"ward_name": "Phường Bến Nghé",
				"ward_code": "HCM-001",
				"ward_type": "Ward",
				"province": "Ho Chi Minh City",
			}
		],
	)
	def test_public_wards_require_and_scope_province(self, get_all, exists, resolve_province):
		result = get_public_wards("HCM")

		self.assertEqual(result["total"], 1)
		self.assertEqual(result["items"][0]["value"], "Ward 1 - Ho Chi Minh City")
		self.assertEqual(result["items"][0]["ward_type"], "Ward")
		resolve_province.assert_called_once_with("HCM")
		exists.assert_called_once_with("CRM Province", "Ho Chi Minh City")
		self.assertEqual(get_all.call_args.kwargs["filters"], {"province": "Ho Chi Minh City"})

	@patch("crm.api.lead_mapping.resolve_ward", return_value="Ward 1 - Ho Chi Minh City")
	@patch("crm.api.lead_mapping.frappe.db.exists", return_value=True)
	@patch(
		"crm.api.lead_mapping.frappe.get_all",
		return_value=[
			{
				"name": "High School 1",
				"school_name": "THPT Nguyễn Huệ",
				"school_code": "HCM-THPT-001",
				"school_type": "Public",
				"province": "Ho Chi Minh City",
				"ward": "Ward 1 - Ho Chi Minh City",
			}
		],
	)
	def test_public_high_schools_scope_ward_and_only_active(self, get_all, exists, resolve_ward):
		result = get_public_high_schools("HCM-001")

		self.assertEqual(result["items"][0]["label"], "THPT Nguyễn Huệ")
		resolve_ward.assert_called_once_with("HCM-001")
		exists.assert_called_once_with("CRM Ward", "Ward 1 - Ho Chi Minh City")
		self.assertEqual(
			get_all.call_args.kwargs["filters"],
			{"ward": "Ward 1 - Ho Chi Minh City", "is_active": 1},
		)

	@patch(
		"crm.api.lead_mapping.frappe.get_all",
		return_value=[
			{
				"name": "Software Engineering",
				"major_name": "Kỹ thuật phần mềm",
				"major_code": "SE",
				"degree_name": "Cử nhân",
				"major_group": "Technology",
			}
		],
	)
	def test_public_majors_return_active_dropdown_items(self, get_all):
		result = get_public_majors()

		self.assertEqual(result["items"][0]["code"], "SE")
		self.assertEqual(get_all.call_args.kwargs["filters"], {"is_active": 1})

	def test_public_leads_filter_campaign_dates_and_paginate(self):
		rows = [
			{
				"name": "LEAD-1",
				"lead_code": "HS-2026-HCM-000001",
				"student_name": "An",
				"phone": "0900000001",
				"email": "an@example.com",
				"major": "Software Engineering",
				"high_school": "SCHOOL-1",
				"province": "PROVINCE-1",
				"ward": "WARD-1",
				"processing_status": "NEW",
				"resolution": "PENDING",
				"campaign": "Campaign 1",
				"creation": "2026-09-15 12:30:00",
			}
		]
		province_rows = [{"name": "PROVINCE-1", "province_name": "Khánh Hòa"}]
		ward_rows = [{"name": "WARD-1", "ward_name": "Khánh Hòa"}]
		school_rows = [{"name": "SCHOOL-1", "school_name": "THPT Nguyễn Văn An"}]
		with (
			patch.object(lead_mapping, "_resolve_campaign_code", return_value="Campaign 1"),
			patch.object(
				lead_mapping.frappe,
				"get_all",
				side_effect=[rows, province_rows, ward_rows, school_rows],
			) as get_all,
			patch.object(lead_mapping.frappe.db, "count", return_value=3) as count,
		):
			result = get_public_leads(
				campaign_code="CAM-2026-00001",
				startdate="2026-09-01",
				enddate="2026-09-30",
				start="10",
				page_length="25",
			)

		expected_leads = [
			{
				**rows[0],
				"high_school": "THPT Nguyễn Văn An",
				"province": "Khánh Hòa",
				"ward": "Khánh Hòa",
			}
		]
		self.assertEqual(result, {"total": 3, "start": 10, "page_length": 25, "leads": expected_leads})
		filters = [
			["campaign", "=", "Campaign 1"],
			["creation", ">=", "2026-09-01 00:00:00"],
			["creation", "<", "2026-10-01 00:00:00"],
		]
		lead_query = get_all.call_args_list[0]
		self.assertEqual(lead_query.kwargs["filters"], filters)
		self.assertEqual(lead_query.kwargs["fields"], list(lead_mapping.PUBLIC_LEAD_LIST_FIELDS))
		self.assertEqual(lead_query.kwargs["limit_start"], 10)
		self.assertEqual(lead_query.kwargs["limit_page_length"], 25)
		count.assert_called_once_with("CRM Lead", filters=filters)

	def test_public_leads_accept_snake_case_date_aliases(self):
		with (
			patch.object(lead_mapping, "_resolve_campaign_code", return_value="Campaign 1"),
			patch.object(lead_mapping.frappe, "get_all", return_value=[]),
			patch.object(lead_mapping.frappe.db, "count", return_value=0),
		):
			result = get_public_leads(
				campaign_code="CAM-2026-00001",
				start_date="2026-09-01",
				end_date="2026-09-30",
			)

		self.assertEqual(result["total"], 0)

	def test_public_leads_reject_invalid_date_range_and_pagination(self):
		with patch.object(lead_mapping, "_resolve_campaign_code", return_value="Campaign 1"):
			with self.assertRaises(LeadMappingError) as context:
				get_public_leads(
					campaign_code="CAM-2026-00001",
					startdate="2026-10-01",
					enddate="2026-09-01",
				)
			self.assertEqual(context.exception.code, "INVALID_DATE")

			with self.assertRaises(LeadMappingError) as context:
				get_public_leads(campaign_code="CAM-2026-00001", page_length=101)
			self.assertEqual(context.exception.code, "INVALID_PAGINATION")

	def test_public_lead_list_is_guest_whitelisted(self):
		source = lead_mapping.__loader__.get_source(lead_mapping.__name__)
		self.assertIn('@frappe.whitelist(allow_guest=True, methods=["GET"])', source)

	def test_internal_batch_payload_allows_missing_campaign_code(self):
		with (
			patch("crm.api.lead_mapping._check_unique"),
			patch("crm.api.lead_mapping._resolve_source", return_value="Promoter"),
			patch("crm.api.lead_mapping._resolve_province", return_value="Hồ Chí Minh"),
			patch("crm.api.lead_mapping.frappe.db.exists", return_value=True),
			patch("crm.api.lead_mapping.frappe.db.get_value", return_value=None),
		):
			values = _normalize_public_lead_payload(
				{
					"student_name": "An",
					"phone": "0981000010",
					"province": "Hồ Chí Minh",
					"source": "Promoter",
				},
				require_campaign=False,
			)

		self.assertIsNone(values["campaign"])


class TestLeadImportCampaignContext(TestCase):
	def test_quick_import_requires_campaign_context(self):
		row = _quick_import_row()
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping, "_parse_import_rows", return_value=[row]),
			patch.object(lead_mapping, "_create_lead") as create_lead,
		):
			with self.assertRaises(LeadMappingError) as context:
				lead_mapping.import_leads.__wrapped__(rows=[row], import_mode="quick_create")

		self.assertEqual(context.exception.code, "CAMPAIGN_REQUIRED")
		create_lead.assert_not_called()

	def test_quick_import_accepts_active_and_closed_campaigns_without_existing_leads(self):
		campaign_code = "CAM-2026-00001"
		for status in ("ACTIVE", "CLOSED"):
			with self.subTest(status=status):
				row = _quick_import_row()
				campaign = SimpleNamespace(name="Campaign 1", status=status)
				with (
					patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
					patch.object(lead_mapping.frappe, "has_permission", return_value=True),
					patch.object(
						lead_mapping.frappe.db,
						"get_value",
						return_value={"name": campaign.name, "status": status},
					) as campaign_lookup,
					patch.object(lead_mapping.frappe.db, "count", return_value=0),
					patch.object(lead_mapping.frappe.db, "savepoint"),
					patch.object(lead_mapping, "_parse_import_rows", return_value=[row]),
					patch.object(
						lead_mapping, "_create_lead", return_value={"name": "LEAD-1"}
					) as create_lead,
				):
					result = lead_mapping.import_leads.__wrapped__(
						rows=[row], import_mode="quick_create", campaign_code=f" {campaign_code.lower()} "
					)

				self.assertEqual(result["created"], 1)
				campaign_lookup.assert_called_once_with(
					"CRM Campaign",
					{"stable_code": campaign_code},
					["name", "status"],
					as_dict=True,
				)
				create_lead.assert_called_once_with(
					row,
					allow_unassigned=True,
					require_import_fields=True,
					campaign_name=campaign.name,
				)

	def test_quick_import_rejects_draft_and_upcoming_campaigns(self):
		row = _quick_import_row()
		for status in ("DRAFT", "UPCOMING"):
			with self.subTest(status=status):
				campaign = SimpleNamespace(name="Campaign 1", status=status)
				with (
					patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
					patch.object(lead_mapping.frappe, "has_permission", return_value=True),
					patch.object(
						lead_mapping.frappe.db,
						"get_value",
						return_value={"name": campaign.name, "status": status},
					),
					patch.object(lead_mapping, "_parse_import_rows", return_value=[row]),
					patch.object(lead_mapping, "_create_lead") as create_lead,
				):
					with self.assertRaises(LeadMappingError) as context:
						lead_mapping.import_leads.__wrapped__(
							rows=[row], import_mode="quick_create", campaign_code="CAM-2026-00001"
						)

				self.assertEqual(context.exception.code, "CAMPAIGN_STATUS_NOT_ALLOWED")
				create_lead.assert_not_called()

	def test_quick_import_requires_campaign_read_permission(self):
		row = _quick_import_row()

		def has_permission(*args, **kwargs):
			doctype = kwargs.get("doctype") or (args[0] if args else None)
			permission = kwargs.get("ptype") or kwargs.get("permission_type")
			if permission is None and len(args) > 1:
				permission = args[1]
			return not (doctype == "CRM Campaign" and permission == "read")

		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(
				lead_mapping.frappe, "has_permission", side_effect=has_permission
			) as permission_check,
			patch.object(
				lead_mapping.frappe.db,
				"get_value",
				return_value={"name": "Campaign 1", "status": "ACTIVE"},
			),
			patch.object(lead_mapping, "_parse_import_rows", return_value=[row]),
			patch.object(lead_mapping, "_create_lead") as create_lead,
		):
			with self.assertRaises(LeadMappingError) as context:
				lead_mapping.import_leads.__wrapped__(
					rows=[row], import_mode="quick_create", campaign_code="CAM-2026-00001"
				)

		self.assertEqual(context.exception.code, "CAMPAIGN_PERMISSION_DENIED")
		self.assertIn(
			("CRM Campaign", "read", "Campaign 1"),
			[call.args for call in permission_check.call_args_list],
		)
		create_lead.assert_not_called()

	def test_quick_import_applies_one_campaign_to_every_row_and_ignores_row_override(self):
		selected_code = "CAM-2026-00001"
		rows = [
			_quick_import_row(student_name="First"),
			_quick_import_row(student_name="Second", campaign_code="CAM-2026-00002"),
		]
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(
				lead_mapping.frappe.db,
				"get_value",
				return_value={"name": "Campaign 1", "status": "ACTIVE"},
			),
			patch.object(lead_mapping.frappe.db, "savepoint"),
			patch.object(lead_mapping, "_parse_import_rows", return_value=rows),
			patch.object(
				lead_mapping,
				"_create_lead",
				side_effect=[{"name": "LEAD-1"}, {"name": "LEAD-2"}],
			) as create_lead,
		):
			result = lead_mapping.import_leads.__wrapped__(
				rows=rows, import_mode="quick_create", campaign_code=selected_code
			)

		self.assertEqual(result["created"], 2)
		created_rows = [call.args[0] for call in create_lead.call_args_list]
		self.assertEqual(
			[row.get("campaign_code") for row in created_rows],
			[None, "CAM-2026-00002"],
		)
		self.assertEqual(
			[call.kwargs["campaign_name"] for call in create_lead.call_args_list],
			["Campaign 1", "Campaign 1"],
		)
		self.assertTrue(all("campaign" not in row for row in created_rows))

	def test_preview_reads_campaign_code_from_multipart_form_dict(self):
		upload = SimpleNamespace(filename="leads.csv", read=lambda: b"ignored")
		row = _quick_import_row(campaign_code="CAM-2026-00002")
		campaign = SimpleNamespace(name="Campaign 1", status="ACTIVE")
		normalized_fields = {
			"student_name": "First",
			"campaign": campaign.name,
			"campaign_code": "CAM-2026-00001",
		}
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
			patch.object(
				lead_mapping.frappe,
				"form_dict",
				{"campaign_code": " CAM-2026-00001 "},
				create=True,
			),
			patch.object(
				lead_mapping,
				"_resolve_quick_import_campaign",
				return_value=campaign.name,
			) as resolve_campaign,
			patch.object(lead_mapping, "_parse_import_file", return_value=[row]),
			patch.object(
				lead_mapping,
				"_normalize_lead_payload",
				return_value=(normalized_fields, [], [], None),
			) as normalize_payload,
			patch.object(lead_mapping, "_create_lead") as create_lead,
		):
			result = lead_mapping.preview_lead_import.__wrapped__()

		self.assertEqual(result["valid"], 1)
		resolve_campaign.assert_called_once_with(" CAM-2026-00001 ")
		self.assertNotIn("campaign_code", normalize_payload.call_args.args[0])
		self.assertNotIn("campaign", result["rows"][0]["fields"])
		self.assertNotIn("campaign_code", result["rows"][0]["fields"])
		create_lead.assert_not_called()

	def test_quick_import_rechecks_campaign_status_after_preview(self):
		upload = SimpleNamespace(filename="leads.csv", read=lambda: b"ignored")
		campaign = SimpleNamespace(name="Campaign 1", status="ACTIVE")
		row = _quick_import_row()
		status = {"value": "ACTIVE"}
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe, "request", SimpleNamespace(files={"file": upload})),
			patch.object(lead_mapping.frappe, "form_dict", {"campaign_code": "CAM-2026-00001"}, create=True),
			patch.object(
				lead_mapping.frappe.db,
				"get_value",
				side_effect=lambda *_args, **_kwargs: {
					"name": campaign.name,
					"status": status["value"],
				},
			),
			patch.object(lead_mapping.frappe.db, "savepoint"),
			patch.object(lead_mapping, "_parse_import_file", return_value=[row]),
			patch.object(lead_mapping, "_parse_import_rows", return_value=[row]),
			patch.object(
				lead_mapping,
				"_normalize_lead_payload",
				return_value=({"student_name": "First"}, [], [], None),
			),
			patch.object(lead_mapping, "_create_lead") as create_lead,
		):
			preview = lead_mapping.preview_lead_import.__wrapped__()
			self.assertEqual(preview["valid"], 1)

			status["value"] = "DRAFT"
			with self.assertRaises(LeadMappingError) as context:
				lead_mapping.import_leads.__wrapped__(
					rows=[row], import_mode="quick_create", campaign_code="CAM-2026-00001"
				)

		self.assertEqual(context.exception.code, "CAMPAIGN_STATUS_NOT_ALLOWED")
		create_lead.assert_not_called()

	def test_legacy_import_does_not_require_campaign_context(self):
		row = _quick_import_row(assigned_to="staff@example.com")
		with (
			patch.object(lead_mapping.frappe, "session", SimpleNamespace(user="Administrator")),
			patch.object(lead_mapping.frappe, "has_permission", return_value=True),
			patch.object(lead_mapping.frappe.db, "savepoint"),
			patch.object(lead_mapping, "_parse_import_rows", return_value=[row]),
			patch.object(lead_mapping, "_create_lead", return_value={"name": "LEAD-1"}) as create_lead,
		):
			result = lead_mapping.import_leads.__wrapped__(rows=[row], import_mode="legacy")

		self.assertEqual(result["created"], 1)
		create_lead.assert_called_once()
		self.assertEqual(create_lead.call_args.args[0], row)
		self.assertFalse(create_lead.call_args.kwargs["allow_unassigned"])
		self.assertFalse(create_lead.call_args.kwargs["require_import_fields"])
		self.assertIsNone(create_lead.call_args.kwargs.get("campaign_name"))


class TestLeadMappingIntegration(FrappeTestCase):
	def test_create_lead_resolves_csv_labels_and_enforces_contract(self):
		from crm.api.lead_mapping import _create_lead

		previous_user = frappe.session.user
		frappe.set_user("Administrator")
		created_name = None
		try:
			result = _create_lead(
				{
					"student_name": "_Lead Mapping API Student",
					"phone": "0981000098",
					"province": "Hồ Chí Minh",
					"source": "Promoter",
					"assigned_to": "ctvsale@gmail.com",
					"advertising_channel": "Facebook",
				}
			)
			created_name = result["name"]
			self.assertEqual(result["doctype"], "CRM Lead")
			self.assertRegex(result["leadCode"], r"^HS-\d{4}-[A-Z]+-\d{6}$")
			self.assertEqual(result["lead_code"], result["leadCode"])
			self.assertIn(
				frappe.db.get_value("CRM Lead", created_name, "province"),
				{"Hồ Chí Minh", "Ho Chi Minh City"},
			)
			self.assertIn(
				frappe.db.get_value("CRM Lead", created_name, "assigned_to"),
				{"CTV Sale", "Cộng tác viên Sale"},
			)
			self.assertEqual(frappe.db.get_value("CRM Lead", created_name, "source"), "Promoter")
			self.assertEqual(
				frappe.db.get_value("CRM Lead", created_name, "advertising_channel"),
				"Facebook",
			)
		finally:
			if created_name and frappe.db.exists("CRM Lead", created_name):
				frappe.delete_doc("CRM Lead", created_name, force=True, ignore_permissions=True)
			frappe.set_user(previous_user)

	def test_import_leads_rolls_back_only_invalid_rows(self):
		from crm.api.lead_mapping import import_leads

		previous_user = frappe.session.user
		frappe.set_user("Administrator")
		created_names = []
		try:
			result = import_leads(
				csv_content=(
					"Họ và Tên,Di động,Tỉnh/Thành Phố,Nguồn,Giao cho\n"
					"_Lead Mapping Import Student,0981000097,Hồ Chí Minh,Promoter,ctvsale@gmail.com\n"
					"_Lead Mapping Import Invalid,not-a-phone,Hồ Chí Minh,Promoter,ctvsale@gmail.com"
				)
			)
			created_names = [student["name"] for student in result["students"]]
			self.assertEqual(result["total"], 2)
			self.assertEqual(result["created"], 1)
			self.assertEqual(result["failed"], 1)
			self.assertEqual(result["errors"][0]["code"], "INVALID_PHONE")
		finally:
			for name in created_names:
				if frappe.db.exists("CRM Lead", name):
					frappe.delete_doc("CRM Lead", name, force=True, ignore_permissions=True)
			frappe.set_user(previous_user)
